"""
Utilities for extracting text from various document formats (PDF, DOCX, etc.)
"""
import os
from typing import Dict, List, Optional, Any

try:
    import PyPDF2
    from docx import Document
except ImportError:
    print("Warning: PyPDF2 or python-docx not installed. Document extraction may fail.")


def extract_text_from_pdf(pdf_path: str) -> str:
    """
    Extract text content from a PDF file.
    
    Args:
        pdf_path: Path to the PDF file
        
    Returns:
        Extracted text content
    """
    try:
        text = []
        with open(pdf_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                text.append(page.extract_text())
        return "\n".join(text)
    except Exception as e:
        print(f"Error extracting text from PDF: {e}")
        return f"Error extracting text from PDF: {e}"


def extract_text_from_docx(docx_path: str) -> str:
    """
    Extract text content from a DOCX file.
    
    Args:
        docx_path: Path to the DOCX file
        
    Returns:
        Extracted text content
    """
    try:
        doc = Document(docx_path)
        paragraphs = [p.text for p in doc.paragraphs]
        return "\n".join(paragraphs)
    except Exception as e:
        print(f"Error extracting text from DOCX: {e}")
        return f"Error extracting text from DOCX: {e}"


def get_document_metadata(file_path: str) -> Dict[str, Any]:
    """
    Extract metadata from a document.
    
    Args:
        file_path: Path to the document file
        
    Returns:
        Dictionary containing document metadata
    """
    metadata = {
        "filename": os.path.basename(file_path),
        "extension": os.path.splitext(file_path)[1],
        "size_bytes": os.path.getsize(file_path),
        "last_modified": os.path.getmtime(file_path)
    }
    
    # Extract format-specific metadata
    if file_path.endswith('.pdf'):
        try:
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                metadata["page_count"] = len(reader.pages)
                if reader.metadata:
                    metadata["title"] = reader.metadata.get('/Title', '')
                    metadata["author"] = reader.metadata.get('/Author', '')
                    metadata["creation_date"] = reader.metadata.get('/CreationDate', '')
        except Exception as e:
            print(f"Error extracting PDF metadata: {e}")
    
    elif file_path.endswith('.docx'):
        try:
            doc = Document(file_path)
            metadata["paragraph_count"] = len(doc.paragraphs)
            core_properties = doc.core_properties
            metadata["title"] = core_properties.title
            metadata["author"] = core_properties.author
            metadata["created"] = core_properties.created
        except Exception as e:
            print(f"Error extracting DOCX metadata: {e}")
    
    return metadata 